"""Create editable Office files from bounded, structured task data on stdin."""
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
import uuid

sys.dont_write_bytecode = True


def text(value, label="text"):
    if not isinstance(value, str):
        raise ValueError(f"{label} must be text")
    return value


def name_for(title):
    title = unicodedata.normalize("NFKC", title)
    title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", title).strip(" .-") or "Office"
    while len(title.encode("utf-8")) > 96:
        title = title[:-1]
    return title + "-" + str(uuid.uuid4())


def image_path(value):
    path = Path(text(value, "image_path"))
    if not path.is_file():
        raise ValueError("image_path must name an existing local image")
    return str(path)


def word(spec, destination):
    from docx import Document
    from docx.shared import Inches, Pt
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = spec.get("font", "Arial")
    normal.font.size = Pt(11)
    from docx.oxml.ns import qn
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), spec.get("font", "Microsoft YaHei"))
    document.add_heading(spec["title"], 0)
    for block in spec["blocks"]:
        kind = block["kind"]
        if kind == "heading":
            document.add_heading(block.get("text", ""), level=block.get("level", 1))
        elif kind in ("paragraph", "bullet"):
            document.add_paragraph(block.get("text", ""), style="List Bullet" if kind == "bullet" else None)
        elif kind == "table":
            rows = block.get("rows", [])
            if not rows or not rows[0]:
                raise ValueError("a table needs at least one nonempty row")
            table = document.add_table(rows=0, cols=max(map(len, rows)))
            table.style = "Table Grid"
            for values in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, values):
                    cell.text = text(value, "table cell")
        elif kind == "image":
            document.add_picture(image_path(block["image_path"]), width=Inches(block.get("width_inches", 5.5)))
        elif kind == "page_break":
            document.add_page_break()
        else:
            raise ValueError("unsupported Word block; use the programmable workflow for advanced edits")
    document.save(destination)
    reopened = Document(destination)
    return {"paragraphs": len(reopened.paragraphs), "tables": len(reopened.tables)}


def presentation(spec, destination):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for entry in spec["slides"]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(.55), Inches(.35), Inches(12.2), Inches(.9))
        title.text_frame.word_wrap = True
        run = title.text_frame.paragraphs[0].add_run()
        run.text = entry["title"]
        run.font.size, run.font.bold = Pt(30), True
        body = entry.get("body", [])
        if body:
            has_media = any(key in entry for key in ("image_path", "chart", "table"))
            box = slide.shapes.add_textbox(Inches(.65), Inches(1.5), Inches(5.4 if has_media else 12), Inches(5.3))
            box.text_frame.word_wrap = True
            for index, value in enumerate(body):
                paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
                paragraph.text = value
                paragraph.font.size = Pt(20)
                paragraph.space_after = Pt(14)
        x, y, width, height = Inches(6.4 if body else .65), Inches(1.65), Inches(6.1 if body else 12), Inches(4.9)
        media = sum(key in entry for key in ("image_path", "chart", "table"))
        if media > 1:
            raise ValueError("one slide supports one image, table, or chart; use the programmable workflow for a custom arrangement")
        if "image_path" in entry:
            from PIL import Image
            with Image.open(image_path(entry["image_path"])) as image:
                ratio = image.width / image.height
            fit_width, fit_height = min(width, height * ratio), min(height, width / ratio)
            slide.shapes.add_picture(image_path(entry["image_path"]), int(x + (width-fit_width)/2), int(y + (height-fit_height)/2), width=int(fit_width), height=int(fit_height))
        if "table" in entry:
            rows = entry["table"]
            if not rows or not rows[0]:
                raise ValueError("a slide table needs nonempty rows")
            table = slide.shapes.add_table(len(rows), max(map(len, rows)), x, y, width, height).table
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        if "chart" in entry:
            chart = entry["chart"]
            if chart["kind"] == "pie" and len(chart["series"]) != 1:
                raise ValueError("a pie chart needs one series; use separate native charts for multiple series")
            data = CategoryChartData()
            data.categories = chart["categories"]
            for series in chart["series"]:
                if len(series["values"]) != len(chart["categories"]):
                    raise ValueError("chart series and categories must have equal lengths")
                data.add_series(series["name"], series["values"])
            kinds = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}
            slide.shapes.add_chart(kinds[chart["kind"]], x, y, width, height, data)
        if "notes" in entry:
            slide.notes_slide.notes_text_frame.text = entry["notes"]
    deck.save(destination)
    reopened = Presentation(destination)
    if len(reopened.slides) != len(spec["slides"]):
        raise ValueError("saved slide count differs from the requested deck")
    return {"slides": len(reopened.slides), "native_objects": sum(len(slide.shapes) for slide in reopened.slides)}


def workbook(spec, destination, max_cells):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils.cell import range_boundaries, coordinate_from_string, column_index_from_string
    book = Workbook()
    book.remove(book.active)
    names = set()
    formulas = 0
    touched = 0
    for entry in spec["sheets"]:
        name = entry["name"]
        if len(name) > 31 or re.search(r'[\\/*?:\[\]]', name) or name.casefold() in names or not name:
            raise ValueError("worksheet names must be unique valid Excel names")
        names.add(name.casefold())
        sheet = book.create_sheet(name)
        if len(entry.get("rows", [])) > 1048576 or any(len(row) > 16384 for row in entry.get("rows", [])):
            raise ValueError("data is outside Excel worksheet coordinates")
        for row in entry.get("rows", []):
            touched += len(row)
            if touched > max_cells:
                raise ValueError("worksheet exceeds its configured cell budget; use the programmable workflow for large local data")
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"  # Data remains text; formulas use the explicit formulas field.
        if entry.get("header", True) and sheet.max_row:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245B8A")
            sheet.freeze_panes = "A2"
        for item in entry.get("column_widths", []):
            if column_index_from_string(item["column"]) > 16384:
                raise ValueError("column width is outside Excel worksheet coordinates")
            sheet.column_dimensions[item["column"]].width = item["width"]
        for formula in entry.get("formulas", []):
            touched += 1
            if touched > max_cells:
                raise ValueError("formulas exceed the configured cell budget")
            column, row = coordinate_from_string(formula["cell"])
            if not 1 <= row <= 1048576 or column_index_from_string(column) > 16384:
                raise ValueError("formula is outside Excel worksheet coordinates")
            expression = formula["expression"]
            if not expression.startswith("="):
                raise ValueError("formula expressions must start with =")
            sheet[formula["cell"]] = expression
            formulas += 1
        for item in entry.get("number_formats", []):
            bounds = range_boundaries(item["range"])
            if any(value is None for value in bounds):
                raise ValueError("number-format ranges must name finite cells")
            if not 1 <= bounds[0] <= bounds[2] <= 16384 or not 1 <= bounds[1] <= bounds[3] <= 1048576:
                raise ValueError("number-format range is outside Excel worksheet coordinates")
            touched += (bounds[2]-bounds[0]+1)*(bounds[3]-bounds[1]+1)
            if touched > max_cells:
                raise ValueError("number-format ranges exceed the configured cell budget")
            for row in sheet.iter_rows(min_col=bounds[0], min_row=bounds[1], max_col=bounds[2], max_row=bounds[3]):
                for cell in row:
                    cell.number_format = item["format"]
    book.save(destination)
    reopened = load_workbook(destination, data_only=False)
    try:
        if reopened.sheetnames != [entry["name"] for entry in spec["sheets"]]:
            raise ValueError("saved worksheet names differ from the request")
        for entry in spec["sheets"]:
            for formula in entry.get("formulas", []):
                if reopened[entry["name"]][formula["cell"]].value != formula["expression"]:
                    raise ValueError("saved formula differs from the request")
    finally:
        reopened.close()
    return {"sheets": len(names), "formulas": formulas}


def recalculate(path, payload, spec, scratch):
    if not any(entry.get("formulas") for entry in spec["sheets"]):
        return path, False
    if not payload.get("node") or not payload.get("cli"):
        raise ValueError("formula recalculation requires the bundled LibreOffice Kit; use the programmable workflow if unavailable")
    output = scratch / "recalculated.xlsx"
    result = subprocess.run([payload["node"], payload["cli"], "recalculate", "--input", str(path), "--output", str(output)], capture_output=True, text=True, timeout=payload["timeout_ms"] / 1000)
    if result.returncode:
        raise ValueError("bundled formula recalculation failed; no calculated values are being claimed")
    from openpyxl import load_workbook
    formulas = load_workbook(output, data_only=False)
    values = load_workbook(output, data_only=True)
    try:
        for entry in spec["sheets"]:
            for item in entry.get("formulas", []):
                if formulas[entry["name"]][item["cell"]].data_type != "f":
                    raise ValueError("recalculation did not preserve an editable formula")
                cell = values[entry["name"]][item["cell"]]
                cached = "" if cell.value is None and cell.data_type in ("s", "str", "inlineStr") else cell.value
                if cached is None or cell.data_type == "e":
                    raise ValueError("formula recalculation returned an unavailable or error result")
                if "expected_value" in item and cached != item["expected_value"]:
                    raise ValueError("calculated formula value differs from the expected result")
    finally:
        formulas.close()
        values.close()
    return output, True


def main():
    started = time.monotonic()
    payload = json.load(sys.stdin)
    spec, kind = payload["spec"], payload["kind"]
    title = text(spec["title"], "title")
    root = Path.cwd()
    scratch = Path(tempfile.mkdtemp(prefix=".office-create-", dir=root))
    os.chmod(scratch, 0o700)
    try:
        source = scratch / ("document." + kind)
        summary = workbook(spec, source, payload["max_cells"]) if kind == "xlsx" else {"docx": word, "pptx": presentation}[kind](spec, source)
        recalculated = False
        if kind == "xlsx":
            source, recalculated = recalculate(source, payload, spec, scratch)
        module_spec = importlib.util.spec_from_file_location("office_structure", Path(__file__).with_name("check_office.py"))
        checker = importlib.util.module_from_spec(module_spec)
        module_spec.loader.exec_module(checker)
        checker.inspect(source)
        destination = root / (name_for(title) + "." + kind)
        # The checked file becomes visible atomically and can never overwrite an existing file.
        os.chmod(source, 0o600)
        os.link(source, destination)
        print(json.dumps({"path": str(destination), "format": kind, "check_summary": json.dumps(summary, ensure_ascii=False), "formula_recalculated": recalculated, "visual_check_required": kind in ("docx", "pptx"), "elapsed_ms": round((time.monotonic()-started)*1000)}, ensure_ascii=False))
    finally:
        shutil.rmtree(scratch)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
