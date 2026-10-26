"""Inspect Office packages and publish SHA-bound edits without rebuilding unrelated parts."""
import copy
import hashlib
import io
import json
import os
from pathlib import Path
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from urllib.parse import unquote

sys.dont_write_bytecode = True
from lxml import etree as ET
from create_office import name_for, configure_standard_streams

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def q(prefix, local):
    return "{" + NS[prefix] + "}" + local


class Package:
    def __init__(self, path, limits):
        suffix = path.suffix.lower()
        if suffix not in (".docx", ".pptx", ".xlsx", ".docm", ".pptm", ".xlsm"):
            raise ValueError("use the programmable Office workflow for legacy or macro-enabled files")
        self.kind = {".docm": "docx", ".pptm": "pptx", ".xlsm": "xlsx"}.get(suffix, suffix[1:])
        with path.open("rb") as stream:
            data = stream.read(limits["max_file_bytes"] + 1)
        if len(data) > limits["max_file_bytes"]:
            raise ValueError("Office source exceeds its configured file budget")
        self.sha = hashlib.sha256(data).hexdigest()
        self.path = path
        self.parts = {}
        self.infos = []
        self.changed = {}
        self.comment = b""
        self.formula_errors = []
        self.formula_error_count = 0
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > limits["max_archive_entries"] or sum(i.file_size for i in infos) > limits["max_uncompressed_bytes"]:
                raise ValueError("Office source exceeds its configured archive budget")
            if len({info.filename for info in infos}) != len(infos):
                raise ValueError("Office source contains duplicate package members")
            for info in infos:
                self.parts[info.filename] = archive.read(info)
                self.infos.append(info)
            self.comment = archive.comment
        if self.kind == "xlsx":
            cells = 0
            for name in self.parts:
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    for node in self.xml(name).iter(q("s", "c")):
                        cells += 1
                        if cells > limits["max_cells"]:
                            raise ValueError("Office source exceeds its configured cell budget; use the programmable workflow for large workbooks")

    def xml(self, part):
        if part not in self.parts:
            raise ValueError("Office source is missing a referenced XML part")
        return ET.fromstring(self.changed.get(part, self.parts[part]), ET.XMLParser(resolve_entities=False, no_network=True))

    def save_xml(self, part, tree):
        self.changed[part] = ET.tostring(tree, encoding="UTF-8", xml_declaration=True)

    def relationships(self, part):
        name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
        if name not in self.parts:
            return {}
        relationships = {}
        for node in self.xml(name):
            if node.get("TargetMode") == "External":
                continue
            target = unquote(node.get("Target", ""))
            target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
            relationships[node.get("Id")] = (node.get("Type", "").rsplit("/", 1)[-1], target)
        return relationships

    def warnings(self):
        warnings = []
        main, namespace = {"docx": ("word/document.xml", "w"), "pptx": ("ppt/presentation.xml", "p"), "xlsx": ("xl/workbook.xml", "s")}[self.kind]
        if self.path.suffix.lower().endswith("m"):
            warnings.append("Macro-enabled documents require the programmable preservation workflow.")
        if not self.xml(main).tag.startswith("{" + NS[namespace] + "}"):
            warnings.append("This OOXML namespace needs the programmable preservation workflow.")
        for name in self.parts:
            low = name.lower()
            if any(token in low for token in ("vbaproject", "activex", "externallinks/", "embeddings/", "diagrams/", "pivot", "slicer", "connections.xml", "querytables/")):
                # Native PowerPoint charts own an ordinary embedded XLSX; preserve it without reading it.
                if not (self.kind == "pptx" and low.startswith("ppt/embeddings/") and low.endswith(".xlsx")):
                    warnings.append("Advanced or active embedded content requires the programmable workflow.")
            if name.endswith(".rels") and any(node.get("TargetMode") == "External" for node in self.xml(name)):
                warnings.append("External relationships require an explicit preservation workflow.")
        if self.kind == "docx":
            tree = self.xml("word/document.xml")
            if tree.xpath(".//w:ins|.//w:del|.//w:sdt|.//w:altChunk|.//w:instrText", namespaces=NS):
                warnings.append("Tracked changes, content controls, or fields require the programmable workflow.")
        return list(dict.fromkeys(warnings))

    def write(self, path):
        with zipfile.ZipFile(path, "x") as archive:
            archive.comment = self.comment
            for info in self.infos:
                archive.writestr(info, self.changed.get(info.filename, self.parts[info.filename]))
        os.chmod(path, 0o600)


def document_parts(package):
    if package.kind == "pptx":
        main = "ppt/presentation.xml"
        relationships = package.relationships(main)
        return [(str(index + 1), relationships[node.get(q("r", "id"))][1])
                for index, node in enumerate(package.xml(main).findall("p:sldIdLst/p:sldId", NS))]
    main = "xl/workbook.xml"
    relationships = package.relationships(main)
    return [(node.get("name"), relationships[node.get(q("r", "id"))][1]) for node in package.xml(main).findall("s:sheets/s:sheet", NS)]


def paragraph_nodes(package):
    return package.xml("word/document.xml").findall(".//w:body//w:p", NS)


def slide_shapes(tree):
    return [node for node in tree.findall(".//p:spTree/p:sp", NS) if node.find("p:txBody", NS) is not None]


def shape_id(node):
    identity = node.find("p:nvSpPr/p:cNvPr", NS)
    return int(identity.get("id"))


def text_of(nodes):
    return "".join(node.text or "" for node in nodes)


def replace_text(nodes, old, new):
    value = text_of(nodes)
    if not isinstance(old, str) or not old or not isinstance(new, str) or value.count(old) != 1:
        raise ValueError("old_text must match exactly once; inspect and select a specific run or shape")
    start, end = value.index(old), value.index(old) + len(old)
    position = 0
    inserted = False
    for node in nodes:
        original = node.text or ""
        left, right = position, position + len(original)
        if right > start and left < end:
            prefix = original[:max(0, start-left)]
            suffix = original[max(0, end-left):] if end < right else ""
            node.text = prefix + (new if not inserted else "") + suffix
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            inserted = True
        position = right


def inspect(package, spec, limits):
    count = min(spec.get("max_items", limits["max_inspect_items"]), limits["max_inspect_items"])
    if count < 1 or any(spec.get(key, 0) < 0 for key in ("paragraph_index", "slide_number")):
        raise ValueError("inspection indices and item counts must be positive (paragraph indices start at zero)")
    preview = limits["max_preview_chars"]
    items = []
    total = 0
    def retain(item):
        nonlocal total
        total += 1
        if total > limits["max_items"]:
            raise ValueError("Office inspection exceeds its configured object budget; select a narrower paragraph, slide, or sheet")
        if len(items) < count:
            items.append(item)
    kind = package.kind
    if kind == "docx":
        for index, paragraph in enumerate(paragraph_nodes(package)):
            if "paragraph_index" in spec and index != spec["paragraph_index"]:
                continue
            runs = paragraph.findall(".//w:r", NS)
            item = {"paragraph_index": index, "text": text_of(paragraph.findall(".//w:t", NS))[:preview], "run_count": len(runs)}
            if "paragraph_index" in spec:
                item["runs"] = [{"run_index": r, "text": text_of(run.findall("w:t", NS))[:preview]} for r, run in enumerate(runs[:count])]
            retain(item)
    elif kind == "pptx":
        for number, part in document_parts(package):
            if "slide_number" in spec and int(number) != spec["slide_number"]:
                continue
            for shape in slide_shapes(package.xml(part)):
                retain({"slide_number": int(number), "shape_id": shape_id(shape), "text": text_of(shape.findall(".//a:t", NS))[:preview]})
    else:
        shared = []
        if "xl/sharedStrings.xml" in package.parts:
            shared = [text_of(si.findall(".//s:t", NS)) for si in package.xml("xl/sharedStrings.xml")]
        for name, part in document_parts(package):
            if "sheet" in spec and spec["sheet"] != name:
                continue
            for cell in package.xml(part).findall("s:sheetData/s:row/s:c", NS):
                value = cell.findtext("s:v", default="", namespaces=NS)
                if cell.get("t") == "s":
                    value = shared[int(value)]
                elif cell.get("t") == "inlineStr":
                    value = text_of(cell.findall("s:is//s:t", NS))
                formula = cell.findtext("s:f", namespaces=NS)
                retain({"sheet": name, "cell": cell.get("r"), "value": value[:preview], "formula": formula[:preview] if formula else formula, "style_id": cell.get("s", "0")})
    summary = {"items": items[:count], "total_items": total, "truncated": total > count}
    while len(json.dumps(summary, ensure_ascii=False).encode("utf-8")) > limits["max_output_bytes"] // 2 and summary["items"]:
        if len(summary["items"]) == 1 and len(summary["items"][0].get("runs", [])) > 1:
            summary["items"][0]["runs"].pop()
        else:
            summary["items"].pop()
        summary["truncated"] = True
    warnings = package.warnings()
    return {"source_sha256": package.sha, "format": package.path.suffix[1:].lower(), "can_patch": not warnings,
            "summary": json.dumps(summary, ensure_ascii=False), "compatibility_warnings": warnings}


def number_format(package, cell, number):
    styles = package.xml("xl/styles.xml")
    formats = styles.find("s:numFmts", NS)
    if formats is None:
        formats = ET.Element(q("s", "numFmts"), count="0")
        styles.insert(0, formats)
    existing = next((node for node in formats if node.get("formatCode") == number), None)
    if existing is None:
        identifier = max([163] + [int(node.get("numFmtId")) for node in formats]) + 1
        ET.SubElement(formats, q("s", "numFmt"), numFmtId=str(identifier), formatCode=number)
        formats.set("count", str(len(formats)))
    else:
        identifier = int(existing.get("numFmtId"))
    xfs = styles.find("s:cellXfs", NS)
    style = copy.deepcopy(xfs[int(cell.get("s", "0"))])
    style.set("numFmtId", str(identifier))
    style.set("applyNumberFormat", "1")
    xfs.append(style)
    xfs.set("count", str(len(xfs)))
    cell.set("s", str(len(xfs)-1))
    package.save_xml("xl/styles.xml", styles)


def cell_coordinates(coordinate):
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    coordinate = coordinate.upper()
    column, row_number = coordinate_from_string(coordinate)
    col_number = column_index_from_string(column)
    if not 1 <= row_number <= 1048576 or col_number > 16384:
        raise ValueError("cell is outside Excel worksheet coordinates")
    return coordinate, col_number, row_number


def worksheet_cell(tree, coordinate):
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    coordinate, col_number, row_number = cell_coordinates(coordinate)
    data = tree.find("s:sheetData", NS)
    row = next((node for node in data if int(node.get("r")) == row_number), None)
    if row is None:
        row = ET.Element(q("s", "row"), r=str(row_number))
        data.append(row)
        data[:] = sorted(data, key=lambda node: int(node.get("r")))
    cell = next((node for node in row if node.get("r") == coordinate), None)
    if cell is None:
        cell = ET.Element(q("s", "c"), r=coordinate)
        inherited = row.get("s")
        if inherited is None:
            inherited = next((node.get("style") for node in tree.findall("s:cols/s:col", NS)
                              if int(node.get("min")) <= col_number <= int(node.get("max")) and node.get("style") is not None), None)
        if inherited is not None:
            cell.set("s", inherited)
        row.append(cell)
        row[:] = sorted(row, key=lambda node: column_index_from_string(coordinate_from_string(node.get("r"))[0]))
    return cell


def range_cells(reference, maximum):
    from openpyxl.utils.cell import get_column_letter, range_boundaries
    left, top, right, bottom = range_boundaries(reference)
    if None in (left, top, right, bottom) or not 1 <= left <= right <= 16384 or not 1 <= top <= bottom <= 1048576:
        raise ValueError("formula result range is outside Excel worksheet coordinates")
    if (right-left+1)*(bottom-top+1) > maximum:
        raise ValueError("formula result range exceeds the configured cell budget")
    return {f"{get_column_letter(column)}{row}" for row in range(top, bottom+1) for column in range(left, right+1)}


def edit_cell(package, tree, operation):
    from openpyxl.utils.cell import range_boundaries
    coordinate, col_number, row_number = cell_coordinates(operation["cell"])
    for merged in tree.findall("s:mergeCells/s:mergeCell", NS):
        left, top, right, bottom = range_boundaries(merged.get("ref"))
        if left <= col_number <= right and top <= row_number <= bottom and (col_number, row_number) != (left, top):
            raise ValueError("a merged cell must be edited through its top-left cell")
    if any(key in operation for key in ("value", "formula")):
        for anchor in tree.findall("s:sheetData/s:row/s:c", NS):
            formula = anchor.find("s:f", NS)
            if formula is not None and formula.get("t") == "array":
                left, top, right, bottom = range_boundaries(formula.get("ref", anchor.get("r")))
                if left <= col_number <= right and top <= row_number <= bottom:
                    raise ValueError("this cell belongs to an array result; edit its inputs or use the programmable workflow to replace the whole formula group")
    cell = worksheet_cell(tree, coordinate)
    original = cell.find("s:f", NS)
    if original is not None and original.get("t") in ("shared", "array", "dataTable") and any(key in operation for key in ("value", "formula")):
        raise ValueError("editing a shared or array formula group requires the programmable workflow")
    if "value" in operation and "formula" in operation:
        raise ValueError("a cell operation must choose value or formula")
    if "value" in operation or "formula" in operation:
        for child in list(cell):
            if child.tag in (q("s", "v"), q("s", "f"), q("s", "is")):
                cell.remove(child)
        cell.attrib.pop("t", None)
        if "formula" in operation:
            expression = operation["formula"]
            if not expression.startswith("="):
                raise ValueError("formula must begin with =")
            ET.SubElement(cell, q("s", "f")).text = expression[1:]
        else:
            value = operation["value"]
            if isinstance(value, str):
                cell.set("t", "inlineStr")
                ET.SubElement(ET.SubElement(cell, q("s", "is")), q("s", "t")).text = value
            elif isinstance(value, bool):
                cell.set("t", "b")
                ET.SubElement(cell, q("s", "v")).text = "1" if value else "0"
            elif value is not None:
                ET.SubElement(cell, q("s", "v")).text = str(value)
    if "number_format" in operation:
        number_format(package, cell, operation["number_format"])
    dimension = tree.find("s:dimension", NS)
    if dimension is not None:
        left, top, right, bottom = range_boundaries(dimension.get("ref", "A1"))
        from openpyxl.utils.cell import get_column_letter
        dimension.set("ref", f"{get_column_letter(min(left, col_number))}{min(top, row_number)}:{get_column_letter(max(right, col_number))}{max(bottom, row_number)}")


def copy_formula_cache(package, fresh, cell, calculated, name, payload):
    """Copy values from the calculation engine without copying its rewritten styles, formulas, or objects."""
    for node in list(cell):
        if node.tag in (q("s", "v"), q("s", "is")):
            cell.remove(node)
    cell.attrib.pop("t", None)
    if calculated is None:
        return
    cached_type = calculated.get("t")
    cached_value = calculated.find("s:v", NS)
    if cached_type == "e":
        package.formula_error_count += 1
        if len(package.formula_errors) < payload["max_inspect_items"]:
            package.formula_errors.append({"sheet": name[:payload["max_preview_chars"]], "cell": cell.get("r"), "value": calculated.findtext("s:v", default="", namespaces=NS)[:payload["max_preview_chars"]]})
    if cached_type in ("s", "inlineStr"):
        if cached_type == "s":
            shared = fresh.xml("xl/sharedStrings.xml")
            cached_text = text_of(shared[int(cached_value.text)].findall(".//s:t", NS))
        else:
            cached_text = text_of(calculated.findall("s:is//s:t", NS))
        cell.set("t", "str")
        ET.SubElement(cell, q("s", "v")).text = cached_text
    else:
        if cached_type is not None:
            cell.set("t", cached_type)
        for node in calculated:
            if node.tag in (q("s", "v"), q("s", "is")):
                cell.append(copy.deepcopy(node))


def merge_formula_caches(package, fresh, tree, calculated_tree, name, payload):
    original = {node.get("r"): node for node in tree.findall("s:sheetData/s:row/s:c", NS)}
    calculated = {node.get("r"): node for node in calculated_tree.findall("s:sheetData/s:row/s:c", NS)}
    formulas = {coordinate for coordinate, node in original.items() if node.find("s:f", NS) is not None}
    current_results, previous_results = set(), set()
    previous_size, current_size = 0, 0
    for coordinate in sorted(formulas):
        formula = original[coordinate].find("s:f", NS)
        computed = calculated.get(coordinate)
        if computed is None or computed.find("s:f", NS) is None:
            raise ValueError("recalculation did not preserve a formula cell")
        if computed.find("s:v", NS) is None and computed.get("t") not in ("str", "inlineStr"):
            raise ValueError("recalculation did not return a usable formula cache")
        next_formula = computed.find("s:f", NS)
        if formula.get("t") == "array":
            members = range_cells(formula.get("ref", coordinate), payload["max_cells"])
            previous_size += len(members)
            if previous_size > payload["max_cells"]:
                raise ValueError("formula caches exceed the configured cell budget")
            previous_results.update(members)
        if next_formula.get("t") == "array":
            reference = next_formula.get("ref", coordinate)
            members = range_cells(reference, payload["max_cells"])
            current_size += len(members)
            if current_size > payload["max_cells"]:
                raise ValueError("formula caches exceed the configured cell budget")
            current_results.update(members)
            formula.set("t", "array")
            formula.set("ref", reference)
        elif formula.get("t") == "array":
            formula.set("ref", coordinate)
    targets = formulas | previous_results | current_results
    if len(targets) > payload["max_cells"]:
        raise ValueError("formula caches exceed the configured cell budget")
    for coordinate in targets:
        value = calculated.get(coordinate) if coordinate in formulas or coordinate in current_results else None
        if coordinate not in original and value is None:
            continue
        cell = original.get(coordinate)
        if cell is None:
            cell = worksheet_cell(tree, coordinate)
        copy_formula_cache(package, fresh, cell, value, name, payload)
    from openpyxl.utils.cell import get_column_letter
    dimension = tree.find("s:dimension", NS)
    if dimension is not None:
        coordinates = [cell_coordinates(coordinate)[1:] for coordinate in original.keys() | targets]
        if coordinates:
            dimension.set("ref", f"{get_column_letter(min(column for column, row in coordinates))}{min(row for column, row in coordinates)}:{get_column_letter(max(column for column, row in coordinates))}{max(row for column, row in coordinates)}")


def calculation_copy(package):
    duplicate = copy.copy(package)
    duplicate.changed = dict(package.changed)
    return duplicate


def calculate_file(source, destination, payload):
    if not payload.get("node") or not payload.get("cli"):
        raise ValueError("genuine formula recalculation requires the bundled LibreOffice Kit")
    process = subprocess.run([payload["node"], payload["cli"], "recalculate", "--input", str(source), "--output", str(destination)], capture_output=True, timeout=payload["timeout_ms"]/1000)
    if process.returncode:
        raise ValueError("formula recalculation failed; original source is unchanged")
    return Package(destination, payload)


def dynamic_formula(package, cell):
    """Only explicit XLDAPR cell metadata permits resizing; ordinary CSE arrays keep their fixed ref."""
    marker = cell.get("cm")
    if marker is None or "xl/metadata.xml" not in package.parts:
        return False
    metadata = package.xml("xl/metadata.xml")
    entries = metadata.findall("s:cellMetadata/s:bk", NS)
    types = metadata.findall("s:metadataTypes/s:metadataType", NS)
    index = int(marker)-1
    if not 0 <= index < len(entries):
        raise ValueError("dynamic formula metadata is incomplete; original source is unchanged")
    for reference in entries[index].findall("s:rc", NS):
        kind = int(reference.get("t"))-1
        if not 0 <= kind < len(types) or types[kind].get("name") != "XLDAPR":
            continue
        future = next((node for node in metadata.findall("s:futureMetadata", NS) if node.get("name") == "XLDAPR"), None)
        records = [] if future is None else future.findall("s:bk", NS)
        entry = int(reference.get("v"))
        if not 0 <= entry < len(records):
            raise ValueError("dynamic formula metadata is incomplete; original source is unchanged")
        return any(node.get("fDynamic") in ("1", "true") for node in records[entry].xpath(".//*[local-name()='dynamicArrayProperties']"))
    return False


def measured_dynamic_range(package, name, part, coordinate, payload, scratch, label):
    from decimal import Decimal
    from openpyxl.utils.cell import get_column_letter
    probe = calculation_copy(package)
    tree = probe.xml(part)
    anchor = worksheet_cell(tree, coordinate)
    formula = anchor.find("s:f", NS)
    if formula is None or not formula.text:
        raise ValueError("dynamic formula cannot be evaluated; use Excel to recalculate the unchanged source")
    expression = formula.text
    if formula.get("ref") is not None:
        cells = {cell.get("r"): cell for cell in tree.findall("s:sheetData/s:row/s:c", NS)}
        for follower in range_cells(formula.get("ref"), payload["max_cells"])-{coordinate}:
            node = cells.get(follower)
            if node is not None:
                for cached in list(node):
                    if cached.tag in (q("s", "v"), q("s", "is")):
                        node.remove(cached)
                node.attrib.pop("t", None)
    formula.attrib.clear()
    # One exact integer carries both dimensions: Excel has at most 16384 columns.
    formula.text = f"ROWS({expression})*16385+COLUMNS({expression})"
    probe.save_xml(part, tree)
    source, destination = scratch / (label+".xlsx"), scratch / (label+"-calculated.xlsx")
    probe.write(source)
    calculated = calculate_file(source, destination, payload)
    cell = worksheet_cell(calculated.xml(dict(document_parts(calculated))[name]), coordinate)
    value = cell.findtext("s:v", namespaces=NS)
    if cell.get("t") not in (None, "n") or value is None:
        raise ValueError("dynamic result size could not be confirmed; use Excel to recalculate the unchanged source")
    number = Decimal(value)
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError("dynamic result size could not be confirmed; original source is unchanged")
    rows, columns = divmod(int(number), 16385)
    if rows < 1 or not 1 <= columns <= 16384:
        raise ValueError("dynamic result size is outside Excel worksheet coordinates")
    _, column, row = cell_coordinates(coordinate)
    reference = f"{coordinate}:{get_column_letter(column+columns-1)}{row+rows-1}"
    range_cells(reference, payload["max_cells"])
    return reference


def prepare_dynamic_calculation(package, payload, scratch):
    targets = []
    for name, part in document_parts(package):
        for cell in package.xml(part).findall("s:sheetData/s:row/s:c", NS):
            if cell.find("s:f", NS) is not None and dynamic_formula(package, cell):
                targets.append((name, part, cell.get("r")))
    calculation = calculation_copy(package)
    measured = []
    for index, (name, part, coordinate) in enumerate(targets):
        reference = measured_dynamic_range(package, name, part, coordinate, payload, scratch, f"size-{index}")
        tree = calculation.xml(part)
        anchor = worksheet_cell(tree, coordinate)
        formula = anchor.find("s:f", NS)
        old = range_cells(formula.get("ref", coordinate), payload["max_cells"])
        new = range_cells(reference, payload["max_cells"])
        cells = {cell.get("r"): cell for cell in tree.findall("s:sheetData/s:row/s:c", NS)}
        for target in new-old:
            existing = cells.get(target)
            if existing is not None and any(node.tag in (q("s", "f"), q("s", "v"), q("s", "is")) for node in existing):
                raise ValueError(f"dynamic result conflicts with an existing cell ({name}!{target}); original source is unchanged")
        for target in old-new:
            existing = cells.get(target)
            if existing is not None:
                for cached in list(existing):
                    if cached.tag in (q("s", "v"), q("s", "is")):
                        existing.remove(cached)
                existing.attrib.pop("t", None)
        formula.set("ref", reference)
        formula.set("t", "array")
        calculation.save_xml(part, tree)
        measured.append((name, part, coordinate, reference))
    return calculation, measured


def patch(package, spec, payload, scratch):
    if package.sha != spec["source_sha256"]:
        raise ValueError("source_sha256 no longer matches; inspect the changed source before editing")
    warnings = package.warnings()
    if warnings:
        raise ValueError("; ".join(warnings) + " Original file is unchanged; continue with the programmable Office workflow.")
    kind = package.kind
    for operation in spec["operations"]:
        if any(operation.get(key, 0) < 0 for key in ("paragraph_index", "run_index", "slide_number", "shape_id")):
            raise ValueError("edit indices cannot be negative")
        if kind == "docx" and operation["kind"] == "word_text":
            part = "word/document.xml"
            tree = package.xml(part)
            paragraph = tree.findall(".//w:body//w:p", NS)[operation["paragraph_index"]]
            if "run_index" in operation:
                paragraph = paragraph.findall(".//w:r", NS)[operation["run_index"]]
            replace_text(paragraph.findall(".//w:t", NS), operation["old_text"], operation["new_text"])
        elif kind == "pptx" and operation["kind"] in ("slide_text", "slide_notes"):
            part = dict(document_parts(package))[str(operation["slide_number"])]
            if operation["kind"] == "slide_notes":
                part = next(target for role, target in package.relationships(part).values() if role == "notesSlide")
                tree = package.xml(part)
                shape = next(node for node in tree.findall(".//p:sp", NS) if node.find("p:nvSpPr/p:nvPr/p:ph", NS) is not None and node.find("p:nvSpPr/p:nvPr/p:ph", NS).get("type") == "body")
            else:
                tree = package.xml(part)
                shape = next(node for node in slide_shapes(tree) if shape_id(node) == operation["shape_id"])
            replace_text(shape.findall(".//a:t", NS), operation["old_text"], operation["new_text"])
        elif kind == "xlsx" and operation["kind"] == "cell":
            part = dict(document_parts(package))[operation["sheet"]]
            tree = package.xml(part)
            edit_cell(package, tree, operation)
        else:
            raise ValueError("operation does not apply to this file; use the programmable workflow for other edits")
        package.save_xml(part, tree)
    output = scratch / ("patched." + kind)
    if kind == "xlsx" and any("value" in item or "formula" in item for item in spec["operations"]) and "xl/calcChain.xml" in package.parts:
        package.infos = [info for info in package.infos if info.filename != "xl/calcChain.xml"]
        types = package.xml("[Content_Types].xml")
        for node in list(types):
            if node.get("PartName") == "/xl/calcChain.xml":
                types.remove(node)
        package.save_xml("[Content_Types].xml", types)
        relations = "xl/_rels/workbook.xml.rels"
        tree = package.xml(relations)
        for node in list(tree):
            if node.get("Type", "").endswith("/calcChain"):
                tree.remove(node)
        package.save_xml(relations, tree)
    calculation, dynamic = prepare_dynamic_calculation(package, payload, scratch) if kind == "xlsx" and any("value" in item or "formula" in item for item in spec["operations"]) else (package, [])
    calculation.write(output)
    recalculated = False
    if kind == "xlsx" and any("value" in item or "formula" in item for item in spec["operations"]):
        formula_parts = [(name, part) for name, part in document_parts(package) if package.xml(part).find(".//s:f", NS) is not None]
        if formula_parts:
            calculated = scratch / "calculated.xlsx"
            fresh = calculate_file(output, calculated, payload)
            fresh_parts = dict(document_parts(fresh))
            for index, (name, part, coordinate, reference) in enumerate(dynamic):
                if measured_dynamic_range(fresh, name, fresh_parts[name], coordinate, payload, scratch, f"confirmed-size-{index}") != reference:
                    raise ValueError("dynamic formula dependencies changed result size; use Excel to recalculate the unchanged source")
            for name, part in formula_parts:
                tree = package.xml(part)
                merge_formula_caches(package, fresh, tree, fresh.xml(fresh_parts[name]), name, payload)
                package.save_xml(part, tree)
            output = scratch / "verified.xlsx"
            package.write(output)
            recalculated = True
            from openpyxl import load_workbook
            values = load_workbook(output, read_only=True, data_only=True)
            try:
                for item in spec["operations"]:
                    if "formula" not in item:
                        continue
                    cell = values[item["sheet"]][item["cell"]]
                    cached = "" if cell.value is None and cell.data_type in ("s", "str", "inlineStr") else cell.value
                    if "expected_value" in item:
                        if cached != item["expected_value"]:
                            raise ValueError("calculated formula value differs from expected_value")
                    elif cached is None or cell.data_type == "e":
                        raise ValueError("a newly entered formula returned an unavailable or error result")
            finally:
                values.close()
    return output, recalculated


def main():
    started = time.monotonic()
    payload = json.load(sys.stdin)
    spec = payload["spec"]
    source = Path(spec["source_path"])
    if source.is_symlink() or not source.is_file():
        raise ValueError("source_path must be a regular Office file")
    package = Package(source, payload)
    if payload["kind"] == "inspect":
        result = inspect(package, spec, payload)
    else:
        scratch = Path(tempfile.mkdtemp(prefix=".office-patch-", dir=Path.cwd()))
        try:
            output, recalculated = patch(package, spec, payload, scratch)
            from check_office import inspect as check_structure
            check_structure(output)
            with source.open("rb") as stream:
                current = stream.read(payload["max_file_bytes"] + 1)
            if hashlib.sha256(current).hexdigest() != package.sha:
                raise ValueError("source changed during editing; original file is preserved and no result was published")
            destination = Path.cwd() / (name_for(spec["title"]) + output.suffix)
            result = {"path": str(destination), "format": output.suffix[1:], "check_summary": json.dumps({"operations": len(spec["operations"]), "source_sha256": package.sha, "changed_parts": sorted(package.changed), "formula_errors": package.formula_errors, "formula_error_count": package.formula_error_count, "formula_errors_truncated": package.formula_error_count > len(package.formula_errors)}),
                      "formula_recalculated": recalculated, "visual_check_required": output.suffix in (".docx", ".pptx"), "elapsed_ms": round((time.monotonic()-started)*1000)}
            if len(json.dumps(result, ensure_ascii=False).encode("utf-8")) > payload["max_output_bytes"]:
                raise ValueError("Office patch result exceeds its configured output budget; narrow the operations or use the programmable workflow")
            os.link(output, destination)
        finally:
            shutil.rmtree(scratch)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    configure_standard_streams()
    try:
        main()
    except Exception as error:
        message = "The selected Office object or required field is unavailable; inspect the source again or use the programmable workflow." if isinstance(error, (IndexError, KeyError, StopIteration)) else str(error)
        print(json.dumps({"error": message}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
